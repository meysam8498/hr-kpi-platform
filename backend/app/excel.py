"""
Excel Export/Import — Generate scoring sheets for managers and import filled results.

Export: Creates an Excel file with team members as rows and KPI criteria as columns.
        Managers fill in scores (0-100) and optionally add comments.

Import: Reads a filled Excel file and updates KPI entries in the database.
"""
from io import BytesIO
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from .models import (
    Employee, Team, KPICriterion, TeamKPIConfig,
    KPIEntry, ReportingPeriod,
)


def export_team_scoring_sheet(
    team_id: int,
    period_id: int,
    db: Session,
) -> BytesIO:
    """
    Generate an Excel scoring sheet for a team in a specific period.

    Layout:
    - Header row: team name, period name, date
    - Column A: employee code
    - Column B: employee name
    - Column C: position
    - Columns D+: one column per KPI criterion (with weight %)
    - Last column: optional comments
    - Empty rows for scores (0-100)
    """
    team = db.query(Team).filter(Team.id == team_id).first()
    period = db.query(ReportingPeriod).filter(ReportingPeriod.id == period_id).first()
    employees = db.query(Employee).filter(Employee.team_id == team_id).order_by(Employee.employee_code).all()

    # Get team KPI criteria with weights
    configs = (
        db.query(TeamKPIConfig)
        .filter(TeamKPIConfig.team_id == team_id, TeamKPIConfig.is_active == True)
        .all()
    )
    criteria = []
    for cfg in configs:
        crit = db.query(KPICriterion).filter(KPICriterion.id == cfg.criterion_id).first()
        if crit:
            criteria.append({"id": crit.id, "name": crit.name, "weight": cfg.weight})

    # Create workbook
    wb = Workbook()
    ws = wb.active
    title = f"{team.name[:15]} - {period.name[:15] if period else 'Scoring'}"
    ws.title = title[:31]

    # Styles
    header_font = Font(name="Tahoma", size=14, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
    subheader_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    subheader_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    data_font = Font(name="Tahoma", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="right", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    score_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + len(criteria))
    title_cell = ws.cell(row=1, column=1, value=f"📋 فرم امتیازدهی KPI — {team.name}")
    title_cell.font = Font(name="Tahoma", size=16, bold=True, color="2B579A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Row 2: Period info
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4 + len(criteria))
    period_cell = ws.cell(row=2, column=1, value=f"دوره: {period.name if period else 'N/A'}  |  تاریخ صدور: {datetime.now().strftime('%Y-%m-%d')}")
    period_cell.font = Font(name="Tahoma", size=11, italic=True, color="666666")
    period_cell.alignment = Alignment(horizontal="center")

    # Row 3: Instructions
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=4 + len(criteria))
    instr_cell = ws.cell(row=3, column=1, value="📌 دستورالعمل: نمرات را در بازه 0 تا 100 وارد کنید. نظرات اختیاری هستند.")
    instr_cell.font = Font(name="Tahoma", size=10, italic=True, color="CC0000")
    instr_cell.alignment = Alignment(horizontal="center")

    # Row 5: Headers
    headers = ["کد پرسنلی", "نام و نام خانوادگی", "سمت"]
    for crit in criteria:
        headers.append(f"{crit['name']}\n(وزن: {crit['weight']}%)")
    headers.append("نظرات")

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.font = subheader_font
        cell.fill = subheader_fill
        cell.alignment = center
        cell.border = thin_border

    ws.row_dimensions[5].height = 50

    # Set column widths
    ws.column_dimensions["A"].width = 15  # employee code
    ws.column_dimensions["B"].width = 25  # name
    ws.column_dimensions["C"].width = 20  # position
    for i, crit in enumerate(criteria):
        col_letter = get_column_letter(4 + i)
        ws.column_dimensions[col_letter].width = 15
    ws.column_dimensions[get_column_letter(4 + len(criteria))].width = 30  # comments

    # Employee rows
    for row_idx, emp in enumerate(employees, 6):
        row_data = [emp.employee_code, emp.full_name, emp.position]
        for crit in criteria:
            # Check if there's an existing score
            existing = (
                db.query(KPIEntry)
                .filter(
                    KPIEntry.employee_id == emp.id,
                    KPIEntry.criterion_id == crit["id"],
                    KPIEntry.period_id == period_id,
                )
                .first()
            )
            row_data.append(existing.score if existing else "")
        row_data.append("")  # comments column

        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            if col_idx <= 3:
                cell.alignment = left
            else:
                cell.alignment = center
                if col_idx >= 4 and col_idx < 4 + len(criteria):
                    cell.fill = score_fill

        ws.row_dimensions[row_idx].height = 30

    # Legend at bottom
    legend_row = len(employees) + 7
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=4 + len(criteria))
    legend = ws.cell(row=legend_row, column=1, value="⚠️ لطفاً فقط سلول‌های خاکستری را پر کنید. نمرات خارج از بازه 0-100 قبول نمی‌شوند.")
    legend.font = Font(name="Tahoma", size=10, color="CC0000")

    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def import_team_scoring_sheet(
    team_id: int,
    period_id: int,
    file_content: bytes,
    db: Session,
) -> dict:
    """
    Import a filled Excel scoring sheet and update KPI entries.

    Returns a summary of what was imported.
    """
    wb = load_workbook(BytesIO(file_content))
    ws = wb.active

    # Get team criteria
    configs = (
        db.query(TeamKPIConfig)
        .filter(TeamKPIConfig.team_id == team_id, TeamKPIConfig.is_active == True)
        .all()
    )
    criteria_map = {}
    for cfg in configs:
        crit = db.query(KPICriterion).filter(KPICriterion.id == cfg.criterion_id).first()
        if crit:
            criteria_map[cfg.criterion_id] = crit.name

    # Parse: rows 6+ are employee data, columns 4+ are criteria
    # Column mapping: col 4 = first criterion, etc.
    crit_ids = list(criteria_map.keys())

    imported = 0
    errors = []
    skipped = []

    for row_idx in range(6, ws.max_row + 1):
        employee_code = ws.cell(row=row_idx, column=1).value
        if not employee_code or str(employee_code).strip() == "":
            continue

        # Find employee
        emp = db.query(Employee).filter(
            Employee.employee_code == str(employee_code).strip(),
            Employee.team_id == team_id,
        ).first()
        if not emp:
            errors.append(f"ردیف {row_idx}: کارمند با کد '{employee_code}' یافت نشد")
            continue

        # Read scores
        for col_offset, crit_id in enumerate(crit_ids):
            col_idx = 4 + col_offset
            score_value = ws.cell(row=row_idx, column=col_idx).value

            if score_value is None or str(score_value).strip() == "":
                skipped.append(f"{emp.full_name} - {criteria_map.get(crit_id, '?')}: خالی")
                continue

            try:
                score = float(score_value)
            except (ValueError, TypeError):
                errors.append(f"{emp.full_name} - {criteria_map.get(crit_id, '?')}: مقدار نامعتبر '{score_value}'")
                continue

            if not (0 <= score <= 100):
                errors.append(f"{emp.full_name} - {criteria_map.get(crit_id, '?')}: نمره {score} خارج از بازه 0-100")
                continue

            # Upsert entry
            existing = (
                db.query(KPIEntry)
                .filter(
                    KPIEntry.employee_id == emp.id,
                    KPIEntry.criterion_id == crit_id,
                    KPIEntry.period_id == period_id,
                )
                .first()
            )
            if existing:
                existing.score = score
            else:
                entry = KPIEntry(
                    employee_id=emp.id,
                    criterion_id=crit_id,
                    period_id=period_id,
                    score=score,
                )
                db.add(entry)

            imported += 1

    db.commit()

    return {
        "imported": imported,
        "skipped": len(skipped),
        "errors": len(errors),
        "error_details": errors,
    }
