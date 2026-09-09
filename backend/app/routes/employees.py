"""
Employee management routes — full CRUD + archive + transfer + bulk import.
"""
from datetime import date
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Employee, Team, User
from ..auth import require_admin_or_hr, get_current_user, visible_team_ids, extra_employee_ids, can_touch_employee
from ..schemas import (EmployeeCreate, EmployeeUpdate, EmployeeOut, BulkImportItem,
                       BulkImportItemByName, BulkImportResult)
from .audit import log_audit

router = APIRouter(prefix="/api/employees", tags=["Employees"])


def _emp_out(emp: Employee, db: Session) -> EmployeeOut:
    team = db.query(Team).filter(Team.id == emp.team_id).first()
    return EmployeeOut(
        id=emp.id, employee_code=emp.employee_code,
        first_name=emp.first_name, last_name=emp.last_name,
        position=emp.position, team_id=emp.team_id,
        hire_date=emp.hire_date, phone=emp.phone,
        is_archived=emp.is_archived,
        created_at=emp.created_at, team_name=team.name if team else None,
    )


@router.get("/", response_model=list[EmployeeOut])
def list_employees(
    team_id: int = None,
    archived: bool = False,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """List employees. By default only active (non-archived).

    Scope: employees and managers see their own team only;
    admin/HR see everyone.
    """
    query = db.query(Employee)
    if user.role == "employee":
        allowed_ids = set([user.employee_id] if user.employee_id else []) | set(extra_employee_ids(user))
        if user.team_id is not None:
            query = query.filter(Employee.team_id == user.team_id)
            if allowed_ids:
                query = query.filter((Employee.team_id == user.team_id) | (Employee.id.in_(list(allowed_ids))))
        elif allowed_ids:
            query = query.filter(Employee.id.in_(list(allowed_ids)))
        else:
            return []
    elif user.role == "manager":
        tids = visible_team_ids(user) or []
        if tids:
            query = query.filter(Employee.team_id.in_(tids))
        else:
            return []
    elif team_id:
        query = query.filter(Employee.team_id == team_id)
    if not archived:
        query = query.filter(Employee.is_archived == False)
    else:
        # If explicitly requesting archived, show only archived
        query = query.filter(Employee.is_archived == True)
    employees = query.order_by(Employee.employee_code).all()
    return [_emp_out(e, db) for e in employees]


@router.get("/all", response_model=list[EmployeeOut])
def list_all_employees(
    team_id: int = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """List ALL employees including archived (same role scoping)."""
    query = db.query(Employee)
    if user.role == "employee":
        allowed_ids = set([user.employee_id] if user.employee_id else []) | set(extra_employee_ids(user))
        if user.team_id is not None:
            query = query.filter((Employee.team_id == user.team_id) | (Employee.id.in_(list(allowed_ids))) if allowed_ids else Employee.team_id == user.team_id)
        elif allowed_ids:
            query = query.filter(Employee.id.in_(list(allowed_ids)))
        else:
            return []
    elif user.role == "manager":
        tids = visible_team_ids(user) or []
        if tids:
            query = query.filter(Employee.team_id.in_(tids))
        else:
            return []
    elif team_id:
        query = query.filter(Employee.team_id == team_id)
    employees = query.order_by(Employee.employee_code).all()
    return [_emp_out(e, db) for e in employees]


@router.get("/import-template")
def download_import_template(_: User = Depends(require_admin_or_hr)):
    """Download the sample Excel template for bulk employee import.
    Same column format as the real HR file: ردیف | شماره پرسنلی | نام |
    نام خانوادگی | واحد | سمت سازمانی | تاریخ استخدام (شمسی) | شماره همراه.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "کارمندان"

    headers = ["ردیف", "شماره پرسنلی", "نام", "نام خانوادگی",
               "واحد", "سمت سازمانی", "تاریخ استخدام", "شماره همراه"]
    header_fill = PatternFill("solid", fgColor="5B6ABF")
    header_font = Font(bold=True, color="FFFFFF", name="Vazirmatn", size=11)

    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    samples = [
        (1, "1002", "پاشا", "بابایی مجید آباد", "فنی", "سرپرست خدمات پس از فروش", "1395/06/01", "09354363073"),
        (2, "1005", "سپیده", "حاجی نوروزی", "فروش", "کارشناس ارشد فروش", "1397/02/01", "09379750607"),
        (3, "1190", "پدرام", "سلیمی", "مالی", "سرپرست واحد مالیاتی", "1405/05/24", "09102003946"),
    ]
    for ri, row in enumerate(samples, 2):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment = Alignment(horizontal="center")

    widths = [7, 14, 14, 22, 14, 34, 15, 16]
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    from fastapi.responses import Response
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=employee_import_template.xlsx"},
    )


@router.get("/{emp_id}", response_model=EmployeeOut)
def get_employee(emp_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    if user.role == "employee":
        allowed = {user.employee_id} if user.employee_id else set()
        allowed.update(extra_employee_ids(user))
        if int(emp_id) not in allowed:
            raise HTTPException(status_code=403, detail="فقط اطلاعات خودتان در دسترس شماست")
    if user.role == "manager" and not can_touch_employee(user, emp):
        raise HTTPException(status_code=403, detail="این کارمند در محدوده دسترسی شما نیست")
    return _emp_out(emp, db)


@router.post("/", response_model=EmployeeOut, status_code=201)
def create_employee(request: EmployeeCreate, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    existing = db.query(Employee).filter(Employee.employee_code == request.employee_code).first()
    if existing:
        raise HTTPException(status_code=400, detail="کد پرسنلی تکراری است")
    team = db.query(Team).filter(Team.id == request.team_id).first()
    if not team:
        raise HTTPException(status_code=400, detail="تیم یافت نشد")
    emp = Employee(
        employee_code=request.employee_code,
        first_name=request.first_name,
        last_name=request.last_name,
        position=request.position,
        team_id=request.team_id,
        hire_date=request.hire_date,
        phone=request.phone,
    )
    db.add(emp)
    db.commit()
    db.refresh(emp)
    log_audit(db, "create", "employee", emp.id, f"ایجاد کارمند «{emp.full_name}» ({emp.employee_code})")
    db.commit()
    return _emp_out(emp, db)


@router.put("/{emp_id}", response_model=EmployeeOut)
def update_employee(emp_id: int, request: EmployeeUpdate, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    update_data = request.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        setattr(emp, field, value)
    db.commit()
    db.refresh(emp)
    log_audit(db, "update", "employee", emp.id, f"ویرایش اطلاعات کارمند «{emp.full_name}»")
    db.commit()
    return _emp_out(emp, db)


@router.put("/{emp_id}/transfer", response_model=EmployeeOut)
def transfer_employee(emp_id: int, new_team_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Transfer an employee to a different team."""
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    team = db.query(Team).filter(Team.id == new_team_id).first()
    if not team:
        raise HTTPException(status_code=400, detail="تیم مقصد یافت نشد")
    old_team = db.query(Team).filter(Team.id == emp.team_id).first()
    emp.team_id = new_team_id
    db.commit()
    db.refresh(emp)
    log_audit(db, "transfer", "employee", emp.id,
              f"انتقال «{emp.full_name}» از تیم «{old_team.name if old_team else '—'}» به تیم «{team.name}»")
    db.commit()
    return _emp_out(emp, db)


@router.put("/{emp_id}/archive", response_model=EmployeeOut)
def archive_employee(emp_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Archive an employee (hide from active views)."""
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    emp.is_archived = True
    db.commit()
    db.refresh(emp)
    log_audit(db, "archive", "employee", emp.id, f"بایگانی کارمند «{emp.full_name}»")
    db.commit()
    return _emp_out(emp, db)


@router.put("/{emp_id}/unarchive", response_model=EmployeeOut)
def unarchive_employee(emp_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Restore an archived employee."""
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    emp.is_archived = False
    db.commit()
    db.refresh(emp)
    log_audit(db, "archive", "employee", emp.id, f"بازگردانی کارمند «{emp.full_name}» از بایگانی")
    db.commit()
    return _emp_out(emp, db)


@router.delete("/{emp_id}")
def delete_employee(emp_id: int, db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Delete an employee and all dependent records (explicit cleanup —
    some child tables lack ORM cascade, e.g. peer_reviews, kpi_results)."""
    emp = db.query(Employee).filter(Employee.id == emp_id).first()
    if not emp:
        raise HTTPException(status_code=404, detail="کارمند یافت نشد")
    name = emp.full_name

    # Explicit cleanup of every table referencing employees.id
    # (order matters only for readability; all are simple FK rows)
    from ..models import (
        KPIEntry, KPIResult, Goal, SelfEvaluation, PIP,
        AbsenceRecord, PeerReview,
    )
    db.query(KPIEntry).filter(KPIEntry.employee_id == emp_id).delete()
    db.query(KPIResult).filter(KPIResult.employee_id == emp_id).delete()
    db.query(Goal).filter(Goal.employee_id == emp_id).delete()
    db.query(SelfEvaluation).filter(SelfEvaluation.employee_id == emp_id).delete()
    db.query(PIP).filter(PIP.employee_id == emp_id).delete()
    db.query(AbsenceRecord).filter(AbsenceRecord.employee_id == emp_id).delete()
    db.query(PeerReview).filter(PeerReview.reviewer_id == emp_id).delete()
    db.query(PeerReview).filter(PeerReview.reviewee_id == emp_id).delete()

    db.delete(emp)
    db.commit()
    log_audit(db, "delete", "employee", emp_id, f"حذف کارمند «{name}» و تمام سوابق مرتبط")
    db.commit()
    return {"message": f"کارمند '{name}' و سوابق مرتبط حذف شد"}


# ──────────────────────────────────────────────
# Bulk Import (JSON or Excel)
# ──────────────────────────────────────────────

@router.post("/import-json", response_model=BulkImportResult)
def bulk_import_json(items: list[BulkImportItem], db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Bulk import employees from a JSON list (team by id)."""
    created = 0
    skipped = []
    errors = []
    for i, item in enumerate(items):
        team = db.query(Team).filter(Team.id == item.team_id).first()
        if not team:
            errors.append({"row": i + 1, "reason": f"تیم با شناسه {item.team_id} یافت نشد"})
            continue
        existing = db.query(Employee).filter(Employee.employee_code == item.employee_code).first()
        if existing:
            skipped.append({"row": i + 1, "employee_code": item.employee_code, "reason": "کد پرسنلی تکراری"})
            continue
        emp = Employee(
            employee_code=item.employee_code, first_name=item.first_name,
            last_name=item.last_name, position=item.position,
            team_id=item.team_id, hire_date=item.hire_date, phone=item.phone,
        )
        db.add(emp)
        created += 1
    db.commit()
    log_audit(db, "import", "employee", None, f"ورود گروهی {created} کارمند جدید")
    db.commit()
    return BulkImportResult(created=created, skipped=skipped, errors=errors)


@router.post("/import-json-by-name", response_model=BulkImportResult)
def bulk_import_json_by_name(items: list[BulkImportItemByName], db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """Bulk import employees from a JSON list — team referenced by *name*;
    missing teams are created automatically. Hire date accepts Jalali
    (yyyy/mm/dd) or Gregorian (yyyy-mm-dd)."""
    from ..utils.jalali import parse_jalali_or_gregorian

    def parse_hire(raw: str):
        return parse_jalali_or_gregorian(raw)

    created = 0
    skipped: list[dict] = []
    errors: list[dict] = []
    teams_cache: dict[str, Team] = {}

    def get_or_create_team(name: str) -> Team | None:
        name = (name or '').strip()
        if not name:
            return None
        if name in teams_cache:
            return teams_cache[name]
        t = db.query(Team).filter(Team.name == name).first()
        if not t:
            t = Team(name=name)
            db.add(t)
            db.flush()
        teams_cache[name] = t
        return t

    for i, item in enumerate(items):
        team = get_or_create_team(item.team_name)
        if not team:
            errors.append({"row": i + 1, "employee_code": item.employee_code, "reason": "نام تیم خالی است"})
            continue
        existing = db.query(Employee).filter(Employee.employee_code == item.employee_code).first()
        if existing:
            skipped.append({"row": i + 1, "employee_code": item.employee_code, "reason": "کد پرسنلی تکراری"})
            continue
        hire = parse_hire(item.hire_date)
        if hire is None:
            errors.append({"row": i + 1, "employee_code": item.employee_code,
                           "reason": f"تاریخ استخدام نامعتبر: {item.hire_date}"})
            continue
        db.add(Employee(
            employee_code=item.employee_code, first_name=item.first_name.strip(),
            last_name=item.last_name.strip(), position=item.position.strip() or "—",
            team_id=team.id, hire_date=hire, phone=item.phone,
        ))
        created += 1
    db.commit()
    log_audit(db, "import", "employee", None,
              f"ورود گروهی {created} کارمند جدید (تیم‌ها بر اساس نام)" +
              (f" — تیم‌های جدید: {', '.join(t.name for t in teams_cache.values() if t.id is None)}" if any(t.id is None for t in teams_cache.values()) else ""))
    db.commit()
    return BulkImportResult(created=created, skipped=skipped, errors=errors)


@router.post("/import-excel", response_model=BulkImportResult)
async def bulk_import_excel(file: UploadFile = File(...), db: Session = Depends(get_db), _: User = Depends(require_admin_or_hr)):
    """
    Bulk import employees from an Excel file — tolerant real-world parser:
      • Header row auto-detected (first non-empty row), columns matched by
        Persian header names, not fixed order.
      • Hire dates accepted as Jalali (yyyy/mm/dd), Gregorian ISO, or Excel
        serial dates — converted to Gregorian ISO.
      • Persian digits in codes/phones converted to Latin.
      • Multiple phone numbers / separators normalized.
      • Unknown teams auto-created.
      • Duplicate codes skipped (never re-created).
    """
    from openpyxl import load_workbook
    from io import BytesIO
    import re

    # ── helpers ──
    PERSIAN_DIGITS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')

    def latin(s: str) -> str:
        return (s or '').translate(PERSIAN_DIGITS)

    def norm(s) -> str:
        """Normalize a cell to a clean string (Arabic ی/ک → Persian, strip)."""
        if s is None:
            return ''
        s = str(s).translate(PERSIAN_DIGITS)
        s = s.replace('\u064a', '\u06cc').replace('\u0643', '\u06a9')  # Arabic → Persian
        return s.strip()

    def parse_hire_date(raw):
        """Accept Jalali yyyy/mm/dd, Gregorian yyyy-mm-dd, datetime, or Excel serial."""
        from datetime import date as _date, datetime as _datetime, timedelta
        if not raw:
            return None
        # openpyxl delivers native datetime/date for date-formatted cells
        if isinstance(raw, _datetime):
            return raw.date()
        if isinstance(raw, _date):
            return raw
        # Excel serial number (date stored as numeric cell)
        if isinstance(raw, (int, float)) and 20000 < raw < 60000:
            return _date(1899, 12, 30) + timedelta(days=int(raw))
        raw = latin(str(raw)).strip().replace('\\', '/').replace('.', '/').replace('-', '/')
        # strip trailing time component if present ("1996/08/22 00/00/00")
        raw = raw.split(' ')[0]
        m = re.match(r'^(\d{4})/(\d{1,2})/(\d{1,2})$', raw)
        if not m:
            # Excel serial number (dates stored as numbers)
            try:
                serial = float(raw)
                if 20000 < serial < 60000:
                    return _date(1899, 12, 30) + timedelta(days=int(serial))
            except ValueError:
                pass
            return None
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if mo > 12 or d > 31:
            return None
        if y > 1700:
            # Gregorian
            try:
                return _date(y, mo, d)
            except ValueError:
                return None
        # Jalali → Gregorian (standard 33-year algorithm, mirrors frontend jalali.ts)
        g = jalali_to_gregorian(y, mo, d)
        return _date(g[0], g[1], g[2])

    def jalali_to_gregorian(jy, jm, jd):
        jy += 1595
        days = -355668 + (365 * jy) + ((jy // 33) * 8) + (((jy % 33) + 3) // 4) + jd
        days += ((jm < 7) and (jm - 1) * 31) or ((jm - 7) * 30 + 186)
        gy = 400 * (days // 146097)
        days %= 146097
        if days > 36524:
            days -= 1
            gy += 100 * (days // 36524)
            days %= 36524
            if days >= 365:
                days += 1
        gy += 4 * (days // 1461)
        days %= 1461
        if days > 365:
            gy += (days - 1) // 365
            days = (days - 1) % 365
        gd = days + 1
        month_days = [0, 31, ((gy % 4 == 0 and gy % 100 != 0) or (gy % 400 == 0)) and 29 or 28,
                      31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
        for gm in range(1, 13):
            if gd <= month_days[gm]:
                break
            gd -= month_days[gm]
        return [gy, gm, gd]

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="فایل اکسل معتبر نیست")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="فایل باید حداقل یک ردیف داده داشته باشد")

    # ── Header detection: find the first row that mentions personnel code ──
    header_idx = 0
    col_map = {}  # canonical → column index
    HEADER_ALIASES = {
        'code': ['شماره پرسنلی', 'کد پرسنلی', 'کد', 'پرسنلی'],
        'first': ['نام', 'نام کوچک'],
        'last': ['نام خانوادگی', 'فامیل'],
        'position': ['سمت', 'سمت سازمانی', 'عنوان شغلی'],
        'team': ['واحد', 'تیم', 'بخش', 'دپارتمان'],
        'hire_date': ['تاریخ استخدام', 'استخدام', 'تاریخ ورود'],
        'phone': ['شماره همراه', 'موبایل', 'تلفن', 'همراه'],
    }
    for i, row in enumerate(rows[:5]):
        cells = [norm(c) for c in row]
        if any(a in c for c in cells for a in HEADER_ALIASES['code'] if c):
            header_idx = i
            for ci, cell in enumerate(cells):
                for key, aliases in HEADER_ALIASES.items():
                    if any(a == cell or (a in cell and len(a) >= 3) for a in aliases):
                        col_map.setdefault(key, ci)
            break
    if 'code' not in col_map or 'first' not in col_map or 'last' not in col_map:
        # Fallback: assume the original fixed order
        col_map = {'code': 1, 'first': 2, 'last': 3, 'position': 5, 'team': 4, 'hire_date': 6, 'phone': 7}

    def cell(row, key):
        ci = col_map.get(key)
        return norm(row[ci]) if ci is not None and ci < len(row) else ''

    created = 0
    skipped = []
    errors = []
    teams_cache: dict[str, int] = {}

    def get_or_create_team(name: str):
        if not name:
            return None
        if name in teams_cache:
            return teams_cache[name]
        t = db.query(Team).filter(Team.name == name).first()
        if not t:
            t = Team(name=name)
            db.add(t)
            db.flush()
        teams_cache[name] = t.id
        return t.id

    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        try:
            code = cell(row, 'code')
            first = cell(row, 'first')
            last = cell(row, 'last')
            position = cell(row, 'position')
            team_name = cell(row, 'team')
            phone_raw = cell(row, 'phone')

            if not (code and first and last):
                errors.append({"row": i, "reason": "نام، نام خانوادگی یا کد پرسنلی خالی است"})
                continue

            hire_date = parse_hire_date(cell(row, 'hire_date'))
            if hire_date is None:
                if (cell(row, 'hire_date') or '').strip() == '':
                    # Empty date in source file → default to start of current Jalali year (placeholder; fix in UI)
                    from ..utils.jalali import jalali_to_gregorian
                    g = jalali_to_gregorian(1405, 1, 1)
                    from datetime import date as _d
                    hire_date = _d(g[0], g[1], g[2])
                else:
                    errors.append({"row": i, "reason": f"تاریخ استخدام نامعتبر: {cell(row, 'hire_date') or '—'}"})
                    continue

            if db.query(Employee).filter(Employee.employee_code == code).first():
                skipped.append({"row": i, "employee_code": code, "reason": "کد پرسنلی تکراری"})
                continue

            # Team created only AFTER duplicate check — re-importing a file
            # must never leave orphan empty teams behind.
            team_id = get_or_create_team(team_name)
            if team_id is None:
                errors.append({"row": i, "reason": "نام واحد/تیم خالی است"})
                continue

            # Normalize phone: keep digits, dash and + only (multi-number cells keep first)
            phone = None
            if phone_raw:
                digits = re.sub(r'[^0-9\-+]', '', latin(phone_raw))
                first_num = digits.split('-')[0] if '-' in digits else digits
                phone = first_num[:20] if first_num else None

            db.add(Employee(
                employee_code=code, first_name=first, last_name=last,
                position=position or "—", team_id=team_id,
                hire_date=hire_date, phone=phone,
            ))
            created += 1
        except Exception as e:
            errors.append({"row": i, "reason": str(e)[:120]})

    db.commit()
    log_audit(db, "import", "employee", None, f"ورود گروهی {created} کارمند از فایل اکسل")
    db.commit()
    return BulkImportResult(created=created, skipped=skipped, errors=errors)

